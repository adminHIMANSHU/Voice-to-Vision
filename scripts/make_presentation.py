from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "VoiceToVision_Presentation.pptx"
IMAGE = ROOT / "data" / "vlm_test.jpg"

W = Inches(13.333)
H = Inches(7.5)
BG = RGBColor(11, 16, 28)
PANEL = RGBColor(22, 30, 48)
PANEL_2 = RGBColor(29, 40, 63)
WHITE = RGBColor(242, 246, 252)
MUTED = RGBColor(170, 184, 207)
CYAN = RGBColor(63, 211, 220)
GREEN = RGBColor(102, 224, 145)
AMBER = RGBColor(255, 194, 77)
RED = RGBColor(255, 102, 112)
PURPLE = RGBColor(171, 137, 255)


def set_fill(shape, color, transparency=0):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.fill.transparency = transparency


def no_line(shape):
    shape.line.fill.background()


def textbox(slide, text, x, y, w, h, size=18, color=WHITE, bold=False,
            font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Pt(2)
    tf.margin_right = Pt(2)
    tf.margin_top = Pt(1)
    tf.margin_bottom = Pt(1)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def rich_text(slide, lines, x, y, w, h, size=16, color=WHITE, bullet=True,
              spacing=5):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Pt(3)
    tf.margin_right = Pt(3)
    for i, item in enumerate(lines):
        if isinstance(item, tuple):
            label, body = item
        else:
            label, body = "", item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(spacing)
        p.level = 0
        if bullet:
            p.text = "• "
        else:
            p.text = ""
        if label:
            r = p.add_run()
            r.text = label
            r.font.bold = True
            r.font.color.rgb = CYAN
            r.font.name = "Aptos"
            r.font.size = Pt(size)
        r = p.add_run()
        r.text = body
        r.font.name = "Aptos"
        r.font.size = Pt(size)
        r.font.color.rgb = color
    return box


def rect(slide, x, y, w, h, color=PANEL, radius=False, transparency=0):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    set_fill(shape, color, transparency)
    no_line(shape)
    return shape


def line(slide, x1, y1, x2, y2, color=CYAN, width=2):
    ln = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = color
    ln.line.width = Pt(width)
    return ln


def base(prs, title, subtitle=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = BG
    rect(slide, 0, 0, 13.333, 0.09, CYAN)
    textbox(slide, title, 0.62, 0.35, 11.9, 0.55, 28, WHITE, True)
    if subtitle:
        textbox(slide, subtitle, 0.65, 0.92, 11.9, 0.35, 11, MUTED)
    return slide


def footer(slide, n):
    textbox(slide, f"VOICETOVISION 2.0  •  ASSISTIVE AI PLATFORM", 0.65, 7.15, 8, 0.2, 8, MUTED, True)
    textbox(slide, f"{n:02d}", 12.2, 7.12, 0.45, 0.22, 10, CYAN, True, align=PP_ALIGN.RIGHT)


def card(slide, x, y, w, h, title, body, accent=CYAN, body_size=14):
    rect(slide, x, y, w, h, PANEL, True)
    rect(slide, x, y, 0.08, h, accent, True)
    textbox(slide, title, x + 0.22, y + 0.18, w - 0.4, 0.32, 16, accent, True)
    textbox(slide, body, x + 0.22, y + 0.62, w - 0.4, h - 0.75, body_size, WHITE)


def add_title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG
    rect(slide, 0, 0, 13.333, 0.12, CYAN)
    rect(slide, 7.75, 0.12, 5.58, 7.38, PANEL, False)
    textbox(slide, "VOICETOVISION", 0.75, 1.0, 6.4, 0.55, 18, CYAN, True)
    textbox(slide, "Multimodal Assistive AI\nfor Independent Mobility", 0.72, 1.68, 6.8, 1.55, 32, WHITE, True)
    textbox(slide, "Real-time spatial navigation • conversational vision • text reading", 0.76, 3.55, 6.25, 0.55, 16, MUTED)
    rect(slide, 0.76, 4.5, 2.1, 0.42, CYAN, True)
    textbox(slide, "PROJECT PRESENTATION", 0.9, 4.59, 1.82, 0.2, 10, BG, True, align=PP_ALIGN.CENTER)
    textbox(slide, "Built with Python, YOLOv11, SmolVLM, Whisper and OpenCV", 0.78, 5.45, 6.0, 0.35, 13, MUTED)
    textbox(slide, "VoiceToVision 2.0", 0.78, 6.58, 3.0, 0.3, 12, GREEN, True)
    if IMAGE.exists():
        pic = slide.shapes.add_picture(str(IMAGE), Inches(8.1), Inches(1.0), width=Inches(4.65), height=Inches(3.3))
        pic.crop_left = 0.04
        pic.crop_right = 0.04
        pic.crop_top = 0.04
        pic.crop_bottom = 0.04
    rect(slide, 8.1, 4.65, 4.65, 1.25, PANEL_2, True)
    textbox(slide, "SEE  •  UNDERSTAND  •  SPEAK", 8.35, 4.93, 4.15, 0.3, 17, CYAN, True, align=PP_ALIGN.CENTER)
    textbox(slide, "An assistive layer between the physical world and the user.", 8.35, 5.38, 4.1, 0.35, 12, WHITE, align=PP_ALIGN.CENTER)
    textbox(slide, "2026", 11.2, 6.72, 1.5, 0.25, 11, MUTED, align=PP_ALIGN.RIGHT)


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    add_title_slide(prs)

    s = base(prs, "The Problem", "Visual information is everywhere, but access is not equal.")
    card(s, 0.7, 1.55, 3.75, 3.95, "Mobility barriers", "Obstacles, people and changing paths are difficult to interpret quickly and safely without continuous assistance.", RED)
    card(s, 4.8, 1.55, 3.75, 3.95, "Information overload", "A camera can see more than a user needs. The system must prioritize the one warning that matters now.", AMBER)
    card(s, 8.9, 1.55, 3.75, 3.95, "Interaction gap", "A useful assistant must listen, understand context and respond in natural language—not only draw boxes.", PURPLE)
    textbox(s, "Design question", 0.75, 6.0, 2, 0.3, 13, CYAN, True)
    textbox(s, "How can a low-cost camera and voice interface become a reliable, context-aware mobility companion?", 2.3, 5.93, 9.9, 0.5, 20, WHITE, True)
    footer(s, 2)

    s = base(prs, "Project Vision", "One platform, three ways to support independence.")
    rect(s, 0.75, 1.45, 11.85, 0.68, PANEL_2, True)
    textbox(s, "Camera + AI perception  →  risk reasoning  →  concise voice guidance", 1.05, 1.64, 11.2, 0.28, 20, WHITE, True, align=PP_ALIGN.CENTER)
    card(s, 0.85, 2.55, 3.65, 2.75, "01  NAVIGATE", "Detect objects, estimate direction and distance, identify approaching hazards, and speak a prioritized action.", CYAN, 15)
    card(s, 4.85, 2.55, 3.65, 2.75, "02  ASK", "Ask natural questions such as “Is my walkway clear?” and receive an answer grounded in the live scene.", GREEN, 15)
    card(s, 8.85, 2.55, 3.65, 2.75, "03  READ", "Point the camera at labels, signs or documents and hear the visible text aloud.", PURPLE, 15)
    textbox(s, "Core principle: assist the user without taking control away from the user.", 1.0, 6.05, 11.3, 0.35, 17, AMBER, True, align=PP_ALIGN.CENTER)
    footer(s, 3)

    s = base(prs, "Objectives & Scope", "Engineering goals translated into measurable system behavior.")
    items = [
        ("Real-time awareness", "Maintain a smooth camera loop while AI tasks run asynchronously."),
        ("Spatial reasoning", "Convert 2D detections into LEFT / CENTER / RIGHT and FAR / MEDIUM / NEAR."),
        ("Safety decisions", "Score hazards using distance, position, confidence and approach motion."),
        ("Natural interaction", "Support Hindi and English voice guidance, questions and reading."),
        ("Actionable UX", "Prefer short, prioritized instructions over noisy streams of detections."),
        ("Deployable foundation", "Run locally with standard camera, microphone and CPU/GPU options."),
    ]
    for i, (t, b) in enumerate(items):
        x = 0.85 + (i % 3) * 4.1
        y = 1.5 + (i // 3) * 2.25
        card(s, x, y, 3.65, 1.7, t, b, [CYAN, GREEN, AMBER, PURPLE, RED, CYAN][i], 13)
    footer(s, 4)

    s = base(prs, "System Architecture", "Independent workers keep perception responsive while richer reasoning runs in the background.")
    nodes = [
        (0.8, 2.1, 1.8, 0.75, "CAMERA\nSTREAM", CYAN),
        (3.1, 1.35, 2.05, 0.75, "YOLOv11\nTRACKER", GREEN),
        (3.1, 3.05, 2.05, 0.75, "WHISPER\nSPEECH", PURPLE),
        (6.0, 1.35, 2.05, 0.75, "SPATIAL\nENGINE", AMBER),
        (6.0, 3.05, 2.05, 0.75, "SMOLVLM\nWORKER", PURPLE),
        (8.9, 2.1, 2.05, 0.75, "SAFETY\nEVALUATOR", RED),
        (11.25, 1.35, 1.35, 0.75, "AUDIO\nQUEUE", GREEN),
        (11.25, 3.05, 1.35, 0.75, "HUD /\nCAPTION", CYAN),
    ]
    for x, y, w, h, label, accent in nodes:
        rect(s, x, y, w, h, PANEL_2, True)
        rect(s, x, y, 0.07, h, accent, True)
        textbox(s, label, x + 0.12, y + 0.18, w - 0.2, h - 0.2, 12, WHITE, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    for a, b in [((2.6, 2.48), (3.1, 1.72)), ((2.6, 2.48), (3.1, 3.42)), ((5.15, 1.72), (6.0, 1.72)), ((5.15, 3.42), (6.0, 3.42)), ((8.05, 1.72), (8.9, 2.48)), ((8.05, 3.42), (8.9, 2.48)), ((10.95, 2.48), (11.25, 1.72)), ((10.95, 2.48), (11.25, 3.42))]:
        line(s, a[0], a[1], b[0], b[1], MUTED, 1.5)
    textbox(s, "Threaded camera, audio and VLM components reduce frame freezes and preserve user feedback.", 1.0, 5.25, 11.3, 0.45, 16, WHITE, True, align=PP_ALIGN.CENTER)
    footer(s, 5)

    s = base(prs, "Perception Layer: YOLOv11 Tracking", "From pixels to persistent object identities.")
    rect(s, 0.75, 1.45, 5.0, 4.9, PANEL, True)
    textbox(s, "Frame-level pipeline", 1.05, 1.76, 4.25, 0.3, 18, CYAN, True)
    steps = ["Capture frame", "Detect objects", "Persist track IDs", "Measure box geometry", "Publish enriched detections"]
    for i, step in enumerate(steps):
        y = 2.35 + i * 0.7
        rect(s, 1.05, y, 0.42, 0.42, CYAN if i < 3 else GREEN, True)
        textbox(s, str(i + 1), 1.05, y + 0.07, 0.42, 0.2, 12, BG, True, align=PP_ALIGN.CENTER)
        textbox(s, step, 1.7, y + 0.04, 3.7, 0.25, 15, WHITE, i == 4)
    rect(s, 6.15, 1.45, 6.45, 4.9, PANEL, True)
    textbox(s, "Detection record", 6.5, 1.76, 5.5, 0.3, 18, GREEN, True)
    record = [
        ("object", "person"),
        ("confidence", "0.92"),
        ("track_id", "17"),
        ("direction", "CENTER"),
        ("distance", "NEAR"),
        ("motion_state", "APPROACHING"),
    ]
    for i, (k, v) in enumerate(record):
        y = 2.35 + i * 0.52
        textbox(s, k, 6.55, y, 2.15, 0.25, 14, MUTED, True)
        textbox(s, v, 8.95, y, 2.7, 0.25, 14, AMBER if k in ("distance", "motion_state") else WHITE, True)
    textbox(s, "Persistent IDs allow the system to reason about change over time—not only isolated frames.", 6.5, 5.55, 5.45, 0.5, 15, WHITE)
    footer(s, 6)

    s = base(prs, "Spatial Reasoning", "A lightweight geometry layer turns detections into navigation language.")
    rect(s, 0.75, 1.45, 7.0, 4.8, PANEL, True)
    textbox(s, "Three-zone orientation", 1.05, 1.75, 3.5, 0.3, 18, CYAN, True)
    rect(s, 1.15, 2.4, 6.2, 2.55, PANEL_2, True)
    line(s, 3.2, 2.4, 3.2, 4.95, CYAN, 1)
    line(s, 5.3, 2.4, 5.3, 4.95, CYAN, 1)
    textbox(s, "LEFT", 1.35, 3.5, 1.55, 0.3, 18, GREEN, True, align=PP_ALIGN.CENTER)
    textbox(s, "CENTER", 3.45, 3.5, 1.6, 0.3, 18, AMBER, True, align=PP_ALIGN.CENTER)
    textbox(s, "RIGHT", 5.45, 3.5, 1.55, 0.3, 18, RED, True, align=PP_ALIGN.CENTER)
    textbox(s, "Object position = horizontal center of bounding box", 1.3, 5.25, 5.8, 0.3, 13, MUTED, align=PP_ALIGN.CENTER)
    rect(s, 8.05, 1.45, 4.55, 4.8, PANEL, True)
    textbox(s, "Proximity + motion", 8.4, 1.75, 3.6, 0.3, 18, AMBER, True)
    rich_text(s, [
        ("NEAR", " large box area"),
        ("MEDIUM", " moderate area"),
        ("FAR", " small area"),
        ("APPROACHING", " positive area change"),
        ("RECEDING", " negative area change"),
    ], 8.4, 2.35, 3.65, 2.2, 15, WHITE, True, 7)
    textbox(s, "Approximate depth is derived from object area ratio; temporal history adds a velocity-like cue.", 8.4, 4.95, 3.55, 0.65, 14, WHITE)
    footer(s, 7)

    s = base(prs, "Safety Evaluation & Risk Logic", "A transparent score makes the action decision explainable.")
    rect(s, 0.8, 1.5, 11.8, 0.9, PANEL_2, True)
    textbox(s, "R = 0.40·distance  +  0.35·position  +  0.25·confidence  +  approach bonus", 1.1, 1.78, 11.2, 0.33, 20, WHITE, True, align=PP_ALIGN.CENTER)
    card(s, 0.85, 2.8, 2.7, 2.1, "DISTANCE", "NEAR = 75\nMEDIUM = 40\nFAR = 10", CYAN, 15)
    card(s, 3.85, 2.8, 2.7, 2.1, "POSITION", "CENTER = 80\nLEFT / RIGHT = 40", AMBER, 15)
    card(s, 6.85, 2.8, 2.7, 2.1, "CONFIDENCE", "Model confidence\n× 100", GREEN, 15)
    card(s, 9.85, 2.8, 2.7, 2.1, "MOTION", "Approaching = +15\nReceding = −10", RED, 15)
    textbox(s, "Hysteresis prevents alert flicker: enter and exit thresholds are intentionally different.", 1.0, 5.55, 11.3, 0.35, 16, WHITE, True, align=PP_ALIGN.CENTER)
    for x, txt, col in [(1.25, "LOW", GREEN), (4.0, "MEDIUM", AMBER), (6.95, "HIGH", RED), (10.0, "CRITICAL", RED)]:
        rect(s, x, 6.05, 1.9, 0.42, col, True)
        textbox(s, txt, x, 6.14, 1.9, 0.18, 11, BG, True, align=PP_ALIGN.CENTER)
    footer(s, 8)

    s = base(prs, "Multimodal Intelligence", "Specialized models answer different questions at the right time.")
    card(s, 0.85, 1.55, 3.65, 3.95, "YOLOv11", "Fast object detection and tracking\n\nBest for:\n• obstacles\n• people and vehicles\n• persistent IDs", GREEN, 15)
    card(s, 4.85, 1.55, 3.65, 3.95, "Whisper tiny", "Speech-to-text for hands-free questions\n\nFlow:\nlisten → transcribe → send question", PURPLE, 15)
    card(s, 8.85, 1.55, 3.65, 3.95, "SmolVLM 256M", "Contextual visual reasoning\n\nBest for:\n• scene summaries\n• visual Q&A\n• text reading", CYAN, 15)
    textbox(s, "The system combines fast, deterministic signals with slower, richer understanding.", 1.0, 6.05, 11.3, 0.35, 17, AMBER, True, align=PP_ALIGN.CENTER)
    footer(s, 9)

    s = base(prs, "Three Operating Modes", "The same camera becomes a navigation tool, an assistant and a reader.")
    modes = [
        ("NAVIGATION", "Continuous obstacle awareness", "Real-time tracking, risk scoring, priority hazard alerts and periodic scene context.", CYAN),
        ("VOICE ASSISTANT", "Ask about the live scene", "Whisper captures a question; SmolVLM answers from the latest camera frame.", GREEN),
        ("TEXT READER", "Make printed information audible", "Read medication labels, books, door numbers, signs and documents.", PURPLE),
    ]
    for i, (title, sub, body, col) in enumerate(modes):
        x = 0.85 + i * 4.1
        rect(s, x, 1.55, 3.65, 4.45, PANEL, True)
        rect(s, x, 1.55, 3.65, 0.62, col, True)
        textbox(s, title, x + 0.15, 1.73, 3.35, 0.24, 14, BG, True, align=PP_ALIGN.CENTER)
        textbox(s, sub, x + 0.3, 2.55, 3.05, 0.55, 18, col, True, align=PP_ALIGN.CENTER)
        textbox(s, body, x + 0.35, 3.45, 2.95, 1.35, 15, WHITE, align=PP_ALIGN.CENTER)
        textbox(s, "TAB to switch", x + 0.4, 5.35, 2.85, 0.25, 11, MUTED, align=PP_ALIGN.CENTER)
    footer(s, 10)

    s = base(prs, "User Interaction Flow", "A concise loop from observation to action.")
    flow = [
        ("1", "Observe", "Camera captures the current scene"),
        ("2", "Interpret", "Detections become spatial facts"),
        ("3", "Evaluate", "Risk engine selects the most urgent hazard"),
        ("4", "Communicate", "Priority speech + subtitle feedback"),
        ("5", "Adapt", "User changes mode, language or asks a question"),
    ]
    for i, (num, title, body) in enumerate(flow):
        x = 0.75 + i * 2.5
        rect(s, x, 2.0, 2.1, 2.3, PANEL, True)
        rect(s, x + 0.7, 1.55, 0.7, 0.7, CYAN if i < 3 else GREEN, True)
        textbox(s, num, x + 0.7, 1.72, 0.7, 0.25, 20, BG, True, align=PP_ALIGN.CENTER)
        textbox(s, title, x + 0.15, 2.6, 1.8, 0.3, 16, WHITE, True, align=PP_ALIGN.CENTER)
        textbox(s, body, x + 0.22, 3.15, 1.65, 0.7, 12, MUTED, align=PP_ALIGN.CENTER)
        if i < 4:
            line(s, x + 2.1, 3.15, x + 2.48, 3.15, CYAN, 2)
    rect(s, 1.2, 5.2, 10.9, 0.72, PANEL_2, True)
    textbox(s, "Example: “Stop! Person directly ahead, approaching.”", 1.45, 5.42, 10.4, 0.25, 18, AMBER, True, align=PP_ALIGN.CENTER)
    footer(s, 11)

    s = base(prs, "Accessible Feedback & HUD", "Every important voice event is also visible as telemetry.")
    card(s, 0.85, 1.5, 3.55, 4.65, "Priority audio queue", "Critical alerts preempt lower-priority messages. Cooldowns reduce repetitive speech and preserve attention.", RED, 15)
    card(s, 4.9, 1.5, 3.55, 4.65, "Live captions", "The latest spoken phrase appears on the HUD, giving caregivers and sighted testers a transparent view of system behavior.", CYAN, 15)
    card(s, 8.95, 1.5, 3.55, 4.65, "Cyber HUD", "Bounding boxes, risk colors, mini-radar, mode, language, FPS and model activity status are rendered in one view.", GREEN, 15)
    textbox(s, "Voice-first design • optional visual telemetry • no vibration hardware required", 0.9, 6.45, 11.5, 0.3, 14, MUTED, True, align=PP_ALIGN.CENTER)
    footer(s, 12)

    s = base(prs, "Technology Stack", "Local-first components selected for an extensible prototype.")
    stack = [
        ("Application", "Python 3.10+", "Orchestration, modes, configuration", CYAN),
        ("Vision", "OpenCV + YOLOv11", "Camera frames, detection, tracking", GREEN),
        ("Reasoning", "SmolVLM 256M", "Scene summaries, Q&A, reading", PURPLE),
        ("Speech", "Whisper + pyttsx3", "Speech recognition and synthesis", AMBER),
        ("Runtime", "Threads + queues", "Non-blocking camera, VLM and audio workers", RED),
        ("Hardware path", "PC / Raspberry Pi / Jetson", "Camera and microphone; haptics configurable", CYAN),
    ]
    for i, (a, b, c, col) in enumerate(stack):
        y = 1.4 + i * 0.8
        rect(s, 0.9, y, 11.6, 0.58, PANEL, True)
        rect(s, 0.9, y, 0.1, 0.58, col, True)
        textbox(s, a, 1.2, y + 0.15, 2.05, 0.22, 14, col, True)
        textbox(s, b, 3.55, y + 0.15, 3.2, 0.22, 14, WHITE, True)
        textbox(s, c, 7.0, y + 0.15, 5.0, 0.22, 13, MUTED)
    footer(s, 13)

    s = base(prs, "Validation & Deployment", "Deterministic checks come before hardware and model-backed smoke tests.")
    card(s, 0.85, 1.5, 3.65, 4.65, "Automated validation", "Spatial zones and distance\nRisk scoring and actions\nHindi alert phrasing\nStale-track cleanup\nPreflight report formatting", GREEN, 15)
    card(s, 4.85, 1.5, 3.65, 4.65, "Operational diagnostics", "scripts/doctor.py checks dependencies and YOLO weights.\n\nOptional camera check validates device access before launch.", CYAN, 15)
    card(s, 8.85, 1.5, 3.65, 4.65, "Launch path", "1. Install requirements\n2. Run preflight\n3. Start with run.bat\n4. Use TAB / SPACE / L / M\n5. Exit with Q or ESC", AMBER, 15)
    textbox(s, "The architecture keeps device-specific integrations configurable rather than hard-coded.", 1.0, 6.45, 11.3, 0.3, 14, MUTED, True, align=PP_ALIGN.CENTER)
    footer(s, 14)

    s = base(prs, "Impact, Roadmap & Conclusion", "VoiceToVision is a foundation for safer, more independent everyday navigation.")
    card(s, 0.85, 1.45, 3.65, 3.8, "Current value", "• Real-time hazard awareness\n• Hands-free visual Q&A\n• Read labels and signs aloud\n• Hindi / English support\n• Explainable risk logic", GREEN, 15)
    card(s, 4.85, 1.45, 3.65, 3.8, "Next steps", "• Field testing with users\n• Better depth estimation\n• Personalised speech profiles\n• Robust low-light handling\n• Optional haptic feedback", PURPLE, 15)
    card(s, 8.85, 1.45, 3.65, 3.8, "Closing message", "See the world.\nUnderstand the scene.\nHear the right action.\n\nTechnology should extend independence—not replace it.", CYAN, 16)
    rect(s, 1.0, 5.75, 11.2, 0.75, PANEL_2, True)
    textbox(s, "Thank you", 1.1, 5.92, 2.2, 0.32, 22, CYAN, True)
    textbox(s, "Questions & demonstration", 8.0, 5.95, 3.75, 0.25, 17, WHITE, True, align=PP_ALIGN.RIGHT)
    footer(s, 15)

    prs.core_properties.title = "VoiceToVision 2.0 Project Presentation"
    prs.core_properties.subject = "Multimodal Assistive AI Platform for the Visually Impaired"
    prs.core_properties.author = "VoiceToVision Project"
    prs.save(OUT)
    print(f"Created {OUT}")


if __name__ == "__main__":
    build()
